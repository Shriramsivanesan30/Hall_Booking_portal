import io
import datetime
from django.shortcuts import render
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.db.models import Count, Sum, Q
from django.utils import timezone
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from apps.bookings.models import Booking
from apps.halls.models import Hall
from apps.departments.models import Department

@login_required
def reports_dashboard_view(request):
    dept_id = request.GET.get('department')
    hall_id = request.GET.get('hall')
    status_filter = request.GET.get('status')
    event_type = request.GET.get('event_type')
    start_date_str = request.GET.get('start_date')
    end_date_str = request.GET.get('end_date')

    today = timezone.localdate()
    start_date = datetime.datetime.strptime(start_date_str, '%Y-%m-%d').date() if start_date_str else today - datetime.timedelta(days=30)
    end_date = datetime.datetime.strptime(end_date_str, '%Y-%m-%d').date() if end_date_str else today + datetime.timedelta(days=30)

    bookings_qs = Booking.objects.select_related('preferred_hall', 'allocated_hall', 'department', 'user')

    bookings_qs = bookings_qs.filter(event_date__gte=start_date, event_date__lte=end_date)

    if dept_id:
        bookings_qs = bookings_qs.filter(department_id=dept_id)
    if hall_id:
        bookings_qs = bookings_qs.filter(Q(allocated_hall_id=hall_id) | Q(preferred_hall_id=hall_id))
    if status_filter:
        bookings_qs = bookings_qs.filter(status=status_filter)
    if event_type:
        bookings_qs = bookings_qs.filter(event_type=event_type)

    # Aggregations
    total_in_range = bookings_qs.count()
    approved_in_range = bookings_qs.filter(status=Booking.STATUS_APPROVED).count()
    pending_in_range = bookings_qs.filter(status__in=[Booking.STATUS_SUBMITTED, Booking.STATUS_HOD_APPROVED, Booking.STATUS_COORDINATOR_APPROVED]).count()
    cancelled_in_range = bookings_qs.filter(status=Booking.STATUS_CANCELLED).count()
    rejected_in_range = bookings_qs.filter(status=Booking.STATUS_REJECTED).count()

    # Department breakdown
    dept_breakdown = Department.objects.filter(is_active=True).annotate(
        count=Count('bookings', filter=Q(bookings__event_date__gte=start_date, bookings__event_date__lte=end_date))
    ).order_by('-count')

    # Hall utilization breakdown
    hall_breakdown = Hall.objects.filter(is_active=True).annotate(
        total_bookings=Count('allocated_bookings', filter=Q(allocated_bookings__event_date__gte=start_date, allocated_bookings__event_date__lte=end_date, allocated_bookings__status=Booking.STATUS_APPROVED)),
        total_hours=Sum('allocated_bookings__duration_hours', filter=Q(allocated_bookings__event_date__gte=start_date, allocated_bookings__event_date__lte=end_date, allocated_bookings__status=Booking.STATUS_APPROVED))
    ).order_by('-total_bookings')

    # Peak hours calculation
    peak_hours_data = [
        {'slot': '09:00 AM - 11:00 AM', 'count': bookings_qs.filter(start_time__lte='10:00', end_time__gte='10:00').count()},
        {'slot': '11:00 AM - 01:00 PM', 'count': bookings_qs.filter(start_time__lte='12:00', end_time__gte='12:00').count()},
        {'slot': '02:00 PM - 04:00 PM', 'count': bookings_qs.filter(start_time__lte='15:00', end_time__gte='15:00').count()},
        {'slot': '04:00 PM - 06:00 PM', 'count': bookings_qs.filter(start_time__lte='17:00', end_time__gte='17:00').count()},
    ]

    departments = Department.objects.filter(is_active=True).order_by('name')
    halls = Hall.objects.filter(is_active=True).order_by('name')

    return render(request, 'reports/reports.html', {
        'bookings': bookings_qs.order_by('-event_date', '-start_time')[:50],
        'total_in_range': total_in_range,
        'approved_in_range': approved_in_range,
        'pending_in_range': pending_in_range,
        'cancelled_in_range': cancelled_in_range,
        'rejected_in_range': rejected_in_range,
        'dept_breakdown': dept_breakdown,
        'hall_breakdown': hall_breakdown,
        'peak_hours_data': peak_hours_data,
        'departments': departments,
        'halls': halls,
        'event_types': [t[0] for t in Booking.EVENT_TYPES],
        'status_choices': Booking.STATUS_CHOICES,
        'start_date': start_date.strftime('%Y-%m-%d'),
        'end_date': end_date.strftime('%Y-%m-%d'),
        'dept_id': dept_id,
        'hall_id': hall_id,
        'status_filter': status_filter,
        'event_type': event_type,
    })

@login_required
def export_excel_report(request):
    """
    Exports filtered bookings report to Microsoft Excel format (.xlsx) using openpyxl.
    """
    dept_id = request.GET.get('department')
    hall_id = request.GET.get('hall')
    status_filter = request.GET.get('status')
    start_date_str = request.GET.get('start_date')
    end_date_str = request.GET.get('end_date')

    bookings_qs = Booking.objects.select_related('preferred_hall', 'allocated_hall', 'department', 'user').order_by('-event_date', '-start_time')

    if start_date_str and end_date_str:
        try:
            start_date = datetime.datetime.strptime(start_date_str, '%Y-%m-%d').date()
            end_date = datetime.datetime.strptime(end_date_str, '%Y-%m-%d').date()
            bookings_qs = bookings_qs.filter(event_date__gte=start_date, event_date__lte=end_date)
        except ValueError:
            pass

    if dept_id:
        bookings_qs = bookings_qs.filter(department_id=dept_id)
    if hall_id:
        bookings_qs = bookings_qs.filter(Q(allocated_hall_id=hall_id) | Q(preferred_hall_id=hall_id))
    if status_filter:
        bookings_qs = bookings_qs.filter(status=status_filter)

    # Create Workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "MCET Seminar Hall Report"

    # Header Styling
    header_fill = PatternFill(start_color="123B72", end_color="123B72", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    border_style = Border(
        left=Side(style='thin', color='D0D7DE'),
        right=Side(style='thin', color='D0D7DE'),
        top=Side(style='thin', color='D0D7DE'),
        bottom=Side(style='thin', color='D0D7DE')
    )

    # Institution Title Row
    ws.merge_cells('A1:J1')
    title_cell = ws['A1']
    title_cell.value = "Dr. Mahalingam College of Engineering and Technology (MCET) - Seminar Hall Booking Report"
    title_cell.font = Font(name="Calibri", size=14, bold=True, color="123B72")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    headers = [
        "Booking ID", "Event Name", "Event Type", "Department",
        "Hall", "Date", "Start Time", "End Time", "Duration (Hrs)", "Status"
    ]

    for col_num, header_title in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_num)
        cell.value = header_title
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border_style

    ws.row_dimensions[3].height = 25

    row_num = 4
    for b in bookings_qs:
        hall = b.get_active_hall()
        row_data = [
            b.booking_id,
            b.event_name,
            b.event_type,
            b.department.code,
            hall.name,
            b.event_date.strftime('%d-%m-%Y'),
            b.start_time.strftime('%I:%M %p'),
            b.end_time.strftime('%I:%M %p'),
            float(b.duration_hours),
            b.get_status_display()
        ]

        for col_num, val in enumerate(row_data, 1):
            cell = ws.cell(row=row_num, column=col_num)
            cell.value = val
            cell.border = border_style
            if col_num in [1, 4, 6, 7, 8, 9, 10]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

        ws.row_dimensions[row_num].height = 20
        row_num += 1

    # Adjust Column widths
    column_widths = [16, 30, 20, 16, 22, 14, 14, 14, 14, 20]
    for i, width in enumerate(column_widths, 1):
        col_letter = openpyxl.utils.get_column_letter(i)
        ws.column_dimensions[col_letter].width = width

    # Save to buffer
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"MCET_Seminar_Hall_Report_{timezone.localdate().strftime('%Y%m%d')}.xlsx"
    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response
